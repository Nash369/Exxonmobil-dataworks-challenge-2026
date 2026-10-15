

from pathlib import Path
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook


# --------------------------------------------------
# FILE LOCATIONS
# --------------------------------------------------

ROOT = Path(__file__).resolve().parent
DOWNLOADS = Path.home() / "Downloads"
PREDICTIONS_PATH = ROOT / "outputs" / "test_predictions.csv"
OUTPUT_PATH = ROOT / "outputs" / "09_Template_Deliverable_Filled.xlsx"

# Assumption: Template Cases 1-10 map to Test Cases 86-95.
CASE_MAPPING = {i: 85 + i for i in range(1, 11)}

TARGET_COLUMNS = {
    "gas_cumulative": "Gas production cumulative [MSCF]",
    "oil_cumulative": "Oil production cumulative [STB]",
    "water_cumulative": "Water production cumulative [STB]",
}


def find_template():
    """Find the original template in Downloads or the project folder."""

    search_locations = [DOWNLOADS, ROOT]

    for folder in search_locations:
        if not folder.exists():
            continue

        matches = sorted(
            folder.glob("09 Template Deliverable*.xlsx"),
            key=lambda path: (
                "(1)" in path.name,
                len(path.name),
            ),
        )

        if matches:
            print(f"Using template: {matches[0]}")
            return matches[0]

    raise FileNotFoundError(
        "Could not find the Excel template.\n"
        "Place '09 Template Deliverable.xlsx' in Downloads or the project folder."
    )


def load_predictions():
    """Load the CSV and standardize its column names."""

    if not PREDICTIONS_PATH.exists():
        raise FileNotFoundError(
            f"Predictions file not found:\n{PREDICTIONS_PATH}\n\n"
            "Run predict_test_cases.py first."
        )

    df = pd.read_csv(PREDICTIONS_PATH)
    df.columns = [str(column).strip() for column in df.columns]

    # Accept either column name used by our prediction scripts.
    if "Prediction" not in df.columns and "Predicted" in df.columns:
        df.rename(columns={"Predicted": "Prediction"}, inplace=True)

    # Repair a malformed first header if necessary.
    if "Case" not in df.columns and "ase" in df.columns:
        df.rename(columns={"ase": "Case"}, inplace=True)

    required = {"Case", "Date", "Target", "Prediction"}
    missing = required - set(df.columns)

    if missing:
        raise ValueError(
            f"Missing CSV columns: {sorted(missing)}\n"
            f"Columns found: {list(df.columns)}"
        )

    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df["Prediction"] = pd.to_numeric(
        df["Prediction"], errors="coerce"
    )

    extracted_case = pd.to_numeric(
        df["Case"].astype(str).str.extract(r"(\d+)")[0],
        errors="coerce",
    )

    if "Case_Number" in df.columns:
        df["Case_Number"] = pd.to_numeric(
            df["Case_Number"], errors="coerce"
        )
        df["Case_Number"] = df["Case_Number"].fillna(extracted_case)
    else:
        df["Case_Number"] = extracted_case

    df.dropna(
        subset=["Date", "Prediction", "Case_Number", "Target"],
        inplace=True,
    )
    df["Case_Number"] = df["Case_Number"].astype(int)

    if df.empty:
        raise ValueError("No valid prediction rows were found in the CSV.")

    print(f"Loaded {len(df):,} prediction rows.")
    return df


def get_curve(df, case_number, target):
    """Get one case's predicted cumulative production curve."""

    curve = df.loc[
        (df["Case_Number"] == case_number)
        & (df["Target"].astype(str) == target),
        ["Date", "Prediction"],
    ].copy()

    curve = (
        curve.dropna()
        .groupby("Date", as_index=False)["Prediction"]
        .mean()
        .sort_values("Date")
    )

    if curve.empty:
        raise ValueError(
            f"No predictions found for Test Case {case_number}, "
            f"target '{target}'."
        )

    # Cumulative production should be non-negative and non-decreasing.
    curve["Prediction"] = (
        curve["Prediction"].clip(lower=0).cummax()
    )

    return curve


def calculate_values(curve, template_dates):
    """Interpolate known predictions and extend beyond the final date."""

    x = np.array(
        [date.toordinal() for date in curve["Date"]],
        dtype=float,
    )
    y = curve["Prediction"].to_numpy(dtype=float)

    target_x = np.array(
        [pd.Timestamp(date).toordinal() for date in template_dates],
        dtype=float,
    )

    # Interpolation within the available prediction period.
    values = np.interp(target_x, x, y)

    # IMPORTANT:
    # The source predictions end around 2008. The extension below is a
    # simple trend assumption, not a validated machine-learning forecast.
    future = target_x > x[-1]

    if future.any():
        start_index = int(
            np.searchsorted(x, x[-1] - 365.25, side="left")
        )

        if start_index >= len(x) - 1:
            start_index = max(0, len(x) - 2)

        elapsed_days = x[-1] - x[start_index]

        if elapsed_days > 0:
            slope = (
                y[-1] - y[start_index]
            ) / elapsed_days
        else:
            slope = 0.0

        # Avoid negative production growth.
        slope = max(0.0, float(slope))

        values[future] = (
            y[-1] + slope * (target_x[future] - x[-1])
        )

    return np.maximum.accumulate(np.maximum(values, 0.0))


def main():
    template_path = find_template()
    predictions = load_predictions()

    workbook = load_workbook(template_path)

    for template_number, test_number in CASE_MAPPING.items():
        sheet_name = f"Case {template_number}"

        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet '{sheet_name}' not found.\n"
                f"Available sheets: {workbook.sheetnames}"
            )

        sheet = workbook[sheet_name]

        headers = {
            str(sheet.cell(row=1, column=column).value).strip(): column
            for column in range(1, sheet.max_column + 1)
            if sheet.cell(row=1, column=column).value is not None
        }

        if "Date" not in headers:
            raise ValueError(f"No Date column found in {sheet_name}.")

        template_rows = []

        for row in range(2, sheet.max_row + 1):
            raw_date = sheet.cell(
                row=row, column=headers["Date"]
            ).value

            if raw_date is None:
                continue

            parsed_date = pd.to_datetime(
                raw_date, dayfirst=True, errors="coerce"
            )

            if pd.notna(parsed_date):
                template_rows.append(
                    (row, parsed_date.to_pydatetime())
                )

        if not template_rows:
            raise ValueError(
                f"No valid dates found in worksheet {sheet_name}."
            )

        template_dates = [date for _, date in template_rows]

        for target, output_column in TARGET_COLUMNS.items():
            if output_column not in headers:
                raise ValueError(
                    f"Column '{output_column}' not found in {sheet_name}."
                )

            curve = get_curve(predictions, test_number, target)

            values = calculate_values(curve, template_dates)

            for (excel_row, _), value in zip(template_rows, values):
                sheet.cell(
                    row=excel_row,
                    column=headers[output_column],
                ).value = float(value)

        print(
            f"Filled {sheet_name} using Test Case {test_number}"
        )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    workbook.properties.description = (
        "Generated using fill_template.py. Mapping Template Cases 1-10 "
        "to Test Cases 86-95 is an assumption. Values beyond the available "
        "prediction period use trend extrapolation and are not validated "
        "machine-learning forecasts."
    )

    workbook.save(OUTPUT_PATH)

    print("\nSUCCESS! Excel workbook created.")
    print(f"Saved to: {OUTPUT_PATH}")
    print(
        "\nReminder: verify the case mapping and future forecast method "
        "before submitting."
    )


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"\nERROR: {error}", file=sys.stderr)
        raise
